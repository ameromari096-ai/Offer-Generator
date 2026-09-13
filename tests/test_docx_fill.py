from pathlib import Path

import docx
import pytest

from offer_agent import fields as fields_module
from offer_agent.docx_fill import fill_placeholders, find_placeholders_in_document

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"


def _full_text(document) -> str:
    """All paragraph text in the document, including inside tables."""
    parts = []

    def _walk(parent):
        for paragraph in parent.paragraphs:
            parts.append(paragraph.text)
        for table in getattr(parent, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    _walk(cell)

    _walk(document)
    return "\n".join(parts)


@pytest.mark.parametrize("template_key", ["purehealth", "talentone"])
def test_template_placeholders_match_known_field_set(template_key):
    template = fields_module.get_template_by_key(template_key)
    found = find_placeholders_in_document(TEMPLATES_DIR / template.filename)
    assert found == set(fields_module.PLACEHOLDERS)


@pytest.mark.parametrize("template_key", ["purehealth", "talentone"])
def test_fill_resolves_all_and_leaves_none_unresolved(template_key, tmp_path):
    template = fields_module.get_template_by_key(template_key)
    values = {name: f"VALUE[{name}]" for name in fields_module.PLACEHOLDERS}

    out_path = tmp_path / "out.docx"
    result = fill_placeholders(TEMPLATES_DIR / template.filename, out_path, values)

    assert not result.unresolved_placeholders
    assert not result.unexpected_placeholders
    assert find_placeholders_in_document(out_path) == set()

    document = docx.Document(str(out_path))
    full_text = _full_text(document)
    assert "VALUE[candidate full name]" in full_text
    assert "VALUE[notice period]" in full_text


def test_fill_retains_unresolved_placeholder_verbatim(tmp_path):
    template = fields_module.get_template_by_key("purehealth")
    values = {name: f"VALUE[{name}]" for name in fields_module.PLACEHOLDERS}
    values["notice period"] = None  # confirmed unavailable

    out_path = tmp_path / "out.docx"
    result = fill_placeholders(TEMPLATES_DIR / template.filename, out_path, values)

    assert result.unresolved_placeholders == ["{{notice period}}"]
    remaining = find_placeholders_in_document(out_path)
    assert remaining == {"notice period"}


def test_branding_and_legal_wording_untouched(tmp_path):
    template = fields_module.get_template_by_key("purehealth")
    values = {name: f"VALUE[{name}]" for name in fields_module.PLACEHOLDERS}
    out_path = tmp_path / "out.docx"
    fill_placeholders(TEMPLATES_DIR / template.filename, out_path, values)

    document = docx.Document(str(out_path))
    full_text = _full_text(document)
    assert "Pure" in full_text and "Health" in full_text
    assert "UAE Labour Law" in full_text
    assert "Senior Manager" in full_text
